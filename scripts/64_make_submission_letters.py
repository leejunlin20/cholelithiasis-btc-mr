#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""64_make_submission_letters.py —— 生成 AJCN 投稿信与利益冲突声明

依 ASN Guide for Authors（2026-03-24 版）
  * Cover Letter：官方"鼓励作者用投稿信明确指出新颖性或其它值得注意之处"
    （"The editors encourage authors to use the cover letter to explicitly indicate
     the novelty or other noteworthy features of their submission."）
    → 本信显式给出 4 条新颖性。
  * Declaration of Interests：ASN 在初投阶段即强制要求提交**每位作者**的利益冲突声明；
    官方另有在线 Declaration of Interests 工具可生成 Word 文件随稿上传。
    → 本脚本产出等价手写件；须由三位作者逐一确认（诚信项，AI 不代作者声明）。

产出（results/submission/letters/）
  Cover_Letter_AJCN.docx
  Declaration_of_Interests.docx
"""
from __future__ import annotations
import os

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(os.environ.get("MR_BASE", "."))
OUT = ROOT / "results" / "submission" / "letters"
OUT.mkdir(parents=True, exist_ok=True)

FONT = "Times New Roman"
TITLE = ("Dietary intake, cholelithiasis and biliary tract cancer: a two-sample "
         "Mendelian randomisation study with cross-ancestry replication and "
         "region-definition-aware colocalisation")
AFFIL = ("Department of Hepatobiliary Surgery, The Central Hospital of Yongzhou, "
         "Yongzhou Hospital Affiliated to University of South China, "
         "Yongzhou, China")
AUTHORS = [("Jun-Lin Li", "0009-0003-4913-9775"),
           ("Yi Jiang", "0009-0002-1161-5230"),
           ("Jinjin Chen", "0009-0001-0564-9218")]


def new_doc(landscape=False):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)
    return d


def _font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)


def para(doc, text="", size=12, bold=False, italic=False, align=None,
         space_after=10, line=1.15, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = line
    if indent is not None:
        pf.left_indent = Inches(indent)
    if text:
        _font(p.add_run(text), size, bold, italic)
    return p


# ============================================================ Cover Letter
def cover_letter():
    d = new_doc()
    para(d, "27 September 2026", 12, space_after=14)

    for line, bold in (("Christopher Duggan, MD, MPH", False),
                       ("Editor-in-Chief", False),
                       ("The American Journal of Clinical Nutrition", False)):
        para(d, line, 12, bold=bold, space_after=0)
    para(d, "", space_after=14)

    para(d, "Dear Dr Duggan,", space_after=12)
    para(d, "We submit for consideration as an Original Research Article the "
            "manuscript entitled:")
    para(d, f"\u201c{TITLE}\u201d", italic=True, indent=0.3, space_after=12)
    para(d, "by Jun-Lin Li, Yi Jiang and Jinjin Chen.")

    para(d, "What is new in this work", 12, bold=True, space_after=4)
    for i, txt in enumerate([
        "This is, to our knowledge, the first causal test of the complete dietary "
        "chain — dietary intake, cholelithiasis and biliary tract cancer — rather "
        "than of a pre-selected handful of nutrients. We screened 379 UK Biobank "
        "diet-related phenotypes against cholelithiasis and carried the analysable "
        "ones through to biliary tract cancer, so that the exposures reported were "
        "selected by a pre-specified rule rather than by prior interest.",

        "Two standard analyses failed systematically in this setting, and we "
        "characterise both rather than only the successful one. Region definition "
        "changed the colocalisation answer (a shared-variant posterior probability "
        "of 0.31 when lead variants were merged within 1 Mb, against 0.19 and 0.12 "
        "when the region was split), and the cholecystitis-adjusted multivariable "
        "model was not identified (conditional F 0.96 and 0.86). Both are silent "
        "failures that would ordinarily be reported as ordinary negative results; "
        "we give reproducible diagnostics for each.",

        "The findings replicate across ancestries. In BioBank Japan the point "
        "estimate for cholelithiasis and biliary tract cancer was essentially "
        "unchanged (odds ratio 1.25 against 1.22 in FinnGen), with concordant "
        "positive and negative control outcomes, and we quantify cross-ancestry "
        "transfer of the genetic instrument directly (regression slope 0.915, "
        "z = 11.17).",

        "We resolve a discrepancy with a recent report rather than leaving it "
        "unaddressed. Jiang et al. (2025) reported shared QTL signals for ABCG5 at "
        "the chromosome 2 locus. We show that the region-level signal there is "
        "prior-dependent, and that the expression-level test does not support a "
        "shared causal variant (posterior probabilities 0.003, 0.031 and 0.241 "
        "across three priors, none above 0.5). Our conclusion is negative at the "
        "expression layer and we say so plainly.",
    ], 1):
        p = para(d, f"{i}. {txt}", space_after=8, indent=0.3)

    para(d, "Points the editors may wish to know at submission", 12, bold=True,
         space_after=4)
    for txt in [
        "This study analysed publicly available summary-level genetic data. It "
        "enrolled no human participants and used no animals, so there is no "
        "participant flowchart in the conventional sense; Panel B of Figure 1 is "
        "the screening flow diagram, from 379 screened dietary phenotypes to 153 "
        "with a genome-wide significant variant, 86 analysable in the full chain "
        "and 38 with at least five instruments.",

        "The study was not prospectively registered. Registration is not "
        "applicable: this is a summary-data Mendelian randomisation analysis "
        "without a systematic review component, so there is no PROSPERO record. "
        "No clinical trial registration number applies.",

        "The STROBE-MR checklist required by the American Society for Nutrition "
        "for Mendelian randomisation studies is uploaded alongside the manuscript, "
        "completed item by item, including the two items we could not satisfy in "
        "full and the reason.",

        "All summary statistics are publicly available (FinnGen release 12, UK "
        "Biobank dietary GWAS via OpenGWAS, BioBank Japan hum0197, eQTL Catalogue "
        "GTEx v8); no ethical approval was required. The manuscript states the "
        "data-sharing position in the required form.",

        "This work received no specific funding. The authors report no conflicts "
        "of interest. Each author's declaration is uploaded with the submission.",
    ]:
        para(d, f"\u2022  {txt}", space_after=8, indent=0.3)

    para(d, "We confirm that this manuscript is original, has not been published "
            "elsewhere, and is not under consideration by any other journal. All "
            "authors have read and approved the submitted version and agree to its "
            "submission to The American Journal of Clinical Nutrition.", space_after=12)

    para(d, "We hope the work is suitable for the Journal and look forward to "
            "your assessment.", space_after=16)
    para(d, "Yours sincerely,", space_after=16)
    para(d, "Jun-Lin Li, MD", 12, bold=True, space_after=0)
    para(d, "Corresponding author", italic=True, space_after=0)
    para(d, AFFIL, space_after=0)
    para(d, "E-mail: leejunlin@live.cn  |  Telephone: +86 746 6236441  |  "
            "ORCID: 0009-0003-4913-9775", space_after=0)
    p = OUT / "Cover_Letter_AJCN.docx"
    d.save(p)
    return p


# ============================================================ Declaration of Interests
def declaration_of_interests():
    d = new_doc()
    para(d, "Declaration of Interests", 14, bold=True, space_after=4)
    para(d, f"Manuscript: \u201c{TITLE}\u201d", italic=True, space_after=10)
    para(d, "Submitted to The American Journal of Clinical Nutrition. The American "
            "Society for Nutrition requires a declaration of interests from every "
            "author at initial submission, including a statement where there is no "
            "conflict.", space_after=12)

    rows = [
        ("Jun-Lin Li", "0009-0003-4913-9775",
         "No conflicts of interest to declare. No financial relationships with any "
         "organisation whose interests could appear to influence the submitted "
         "work; no patents, no stock or share ownership, no honoraria, no travel "
         "grants and no advisory roles relevant to this work."),
        ("Yi Jiang", "0009-0002-1161-5230",
         "No conflicts of interest to declare."),
        ("Jinjin Chen", "0009-0001-0564-9218",
         "No conflicts of interest to declare."),
    ]
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(["Author", "ORCID", "Declaration",
                           "Signature / date"]):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(h), 10.5, bold=True)
    for a, o, decl in rows:
        cells = t.add_row().cells
        for i, v in enumerate([a, o, decl, ""]):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run(v), 10.5)

    para(d, "", space_after=10)
    para(d, "Note for the corresponding author", 12, bold=True, space_after=4)
    para(d, "The American Society for Nutrition also provides an online Declaration "
            "of Interests tool that generates a Word file for upload with the "
            "submission. If you generate that file, upload it in place of this one; "
            "the two are equivalent in content. Every author must confirm their own "
            "declaration before the file is uploaded — the text above reflects the "
            "statement already made in the manuscript and must not be treated as "
            "author-confirmed until each author has seen it.", italic=True,
         space_after=10)
    p = OUT / "Declaration_of_Interests.docx"
    d.save(p)
    return p


if __name__ == "__main__":
    for f in (cover_letter(), declaration_of_interests()):
        print(f"[写出] {f}  ({f.stat().st_size/1024:.0f} KB)")
